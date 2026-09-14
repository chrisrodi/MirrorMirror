# mirror_mirror.py
# Captures exactly the glass region using Qt grabWindow (logical coords).
# Glass pane blocks clicks, can be dragged from center or border, and resized via edges/corners.
# Pauses when adjusting/moving/resizing or when OCR text hasn't changed.
# Answers are shown in a liquid-glass card INSIDE the pane. The card's rectangle is
# blanked out of every screenshot before OCR, so the answer text can never feed back
# into the model.
# deps: pyqt6 pillow openai python-dotenv pyobjc-framework-Cocoa pyobjc-framework-Vision
#       pypdf python-docx python-pptx   (reference-file text extraction)
#       (pytesseract + a tesseract binary are an optional fallback OCR backend)
# model: gpt-6-astra by default (see MM_MODEL / MM_REASONING_EFFORT)
import os, sys, time, threading, shutil, json, re
from pathlib import Path
from dataclasses import dataclass, field
from typing import Callable, List, Optional, Tuple
import uuid

from dotenv import load_dotenv
dotenv_path = Path(__file__).with_name(".env")
load_dotenv(dotenv_path=dotenv_path)

from PyQt6.QtCore import (
    Qt, QRect, QRectF, QPoint, QTimer, pyqtSignal, QObject, pyqtSlot, QStandardPaths,
    QBuffer, QIODevice
)
from PyQt6.QtGui import (
    QFont, QColor, QPainter, QPen, QBrush, QGuiApplication, QImage, QRegion,
    QLinearGradient, QPainterPath, QFontMetrics
)
from PyQt6.QtWidgets import (
    QSpacerItem, QSizePolicy, QColorDialog, QCheckBox, QWidget, QApplication, QLabel,
    QFrame, QVBoxLayout, QHBoxLayout, QPushButton, QDialog, QLineEdit, QScrollArea,
    QFileDialog
)
APP_NAME = "MirrorMirror"
APP_VERSION = "0.1.0"

# ---------- OCR backends ----------
# Default: Apple's Vision framework (ships with macOS, nothing to install or bundle).
# Fallback: Tesseract via pytesseract, only if both are present on the machine.
from PIL import Image, ImageOps, ImageFilter

try:
    import objc
    import Vision
    from Foundation import NSData, NSURL
    _HAVE_VISION = sys.platform == "darwin"
except Exception:
    _HAVE_VISION = False

try:
    import pytesseract
    for _cand in ("/opt/homebrew/bin/tesseract", "/usr/local/bin/tesseract"):
        if os.path.exists(_cand):
            pytesseract.pytesseract.tesseract_cmd = _cand
            break
    _TESS = shutil.which("tesseract") or pytesseract.pytesseract.tesseract_cmd
    pytesseract.pytesseract.tesseract_cmd = _TESS
    _HAVE_TESS = os.path.exists(_TESS)
except Exception:
    pytesseract = None
    _HAVE_TESS = False

_ocr_env = (os.getenv("MM_OCR") or "").strip().lower()
if _ocr_env in ("vision", "tesseract"):
    OCR_BACKEND = _ocr_env
else:
    OCR_BACKEND = "vision" if _HAVE_VISION else ("tesseract" if _HAVE_TESS else "none")


def _vision_recognize(handler) -> str:
    """Run VNRecognizeTextRequest on a prepared handler. Returns lines top-to-bottom."""
    with objc.autorelease_pool():
        req = Vision.VNRecognizeTextRequest.alloc().init()
        req.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
        req.setUsesLanguageCorrection_(True)
        ok, err = handler.performRequests_error_([req], None)
        if not ok:
            raise RuntimeError(str(err) if err is not None else "Vision request failed")
        rows = []
        for obs in (req.results() or []):
            cands = obs.topCandidates_(1)
            if not cands:
                continue
            bb = obs.boundingBox()  # normalized, origin bottom-left
            top = bb.origin.y + bb.size.height
            rows.append((round(-top * 50), bb.origin.x, str(cands[0].string())))
        rows.sort()
        return "\n".join(t for _, _, t in rows).strip()


def ocr_vision(qimg) -> str:
    """OCR a QImage (screen capture) with Apple Vision."""
    buf = QBuffer()
    buf.open(QIODevice.OpenModeFlag.WriteOnly)
    qimg.save(buf, "PNG")
    ba = buf.data()
    data = NSData.dataWithBytes_length_(ba.data(), ba.size())
    # options must be None here; pyobjc rejects an empty dict for this initializer
    handler = Vision.VNImageRequestHandler.alloc().initWithData_options_(data, None)
    return _vision_recognize(handler)


def ocr_image_file(path: str) -> str:
    """OCR an image file (png/jpg/heic/tiff/…) with Apple Vision, straight from disk."""
    if not _HAVE_VISION:
        raise RuntimeError("Image references need Apple Vision (macOS)")
    url = NSURL.fileURLWithPath_(str(path))
    handler = Vision.VNImageRequestHandler.alloc().initWithURL_options_(url, None)
    return _vision_recognize(handler)


def ocr_tesseract(qimg) -> str:
    if not _HAVE_TESS:
        raise RuntimeError("Tesseract is not installed")
    img = pil_from_qimage(qimg).convert("L")
    img = ImageOps.autocontrast(img, cutoff=2)
    img = img.filter(ImageFilter.MedianFilter(3))
    return pytesseract.image_to_string(img, config="--oem 3 --psm 6 -l eng").strip()


def run_ocr(qimg) -> str:
    if OCR_BACKEND == "vision":
        return ocr_vision(qimg)
    if OCR_BACKEND == "tesseract":
        return ocr_tesseract(qimg)
    raise RuntimeError("No OCR backend available (Vision needs macOS; or install tesseract)")

# ---------- OpenAI -----------
from openai import OpenAI
# Default to OpenAI's current flagship. Override with MM_MODEL if you want
# something cheaper/faster, e.g. MM_MODEL=gpt-5.6-terra.
DEFAULT_MODEL = "gpt-6-astra"
MODEL = os.getenv("MM_MODEL") or DEFAULT_MODEL
# Reasoning models take an effort level instead of temperature.
# low = fastest/cheapest, medium = balanced, high/xhigh/max = slower, deeper.
REASONING_EFFORT = os.getenv("MM_REASONING_EFFORT", "medium")

# ---------- API key storage ----------
# The key lives in a per-user config file outside the project, so nothing
# secret has to be checked into git or bundled into the app.

def config_path() -> Path:
    base = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.GenericConfigLocation)
    if not base:
        base = str(Path.home() / ".config")
    return Path(base) / APP_NAME / "config.json"

def load_config() -> dict:
    try:
        with open(config_path(), "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return {}

def save_config(data: dict) -> None:
    path = config_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".json.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    try:
        os.chmod(tmp, 0o600)  # owner read/write only
    except OSError:
        pass
    os.replace(tmp, path)

def get_api_key() -> str:
    """Stored key wins; OPENAI_API_KEY from the environment/.env is a fallback."""
    key = (load_config().get("openai_api_key") or "").strip()
    if key:
        return key
    return (os.getenv("OPENAI_API_KEY") or "").strip()

def set_api_key(key: str) -> None:
    global _client
    cfg = load_config()
    cfg["openai_api_key"] = key.strip()
    save_config(cfg)
    _client = None  # force re-creation with the new key

def mask_key(key: str) -> str:
    if not key:
        return "not set"
    if len(key) <= 8:
        return "•" * len(key)
    return f"{key[:3]}…{key[-4:]}"

_client = None
_client_lock = threading.Lock()

def get_client() -> OpenAI:
    """Lazily build the OpenAI client so importing the module never needs a key."""
    global _client
    with _client_lock:
        if _client is None:
            key = get_api_key()
            if not key:
                raise RuntimeError("No OpenAI API key configured.")
            _client = OpenAI(api_key=key)
        return _client

def verify_api_key(key: str) -> str:
    """Return '' if the key works, otherwise a short error message."""
    try:
        OpenAI(api_key=key, timeout=15.0, max_retries=0).models.list()
        return ""
    except Exception as e:
        msg = str(e)
        if "401" in msg or "invalid_api_key" in msg or "Incorrect API key" in msg:
            return "OpenAI rejected this key. Check it and try again."
        return f"Could not verify the key: {msg}"


# ================== reference files ==================
# Files the user adds are read ON THIS MAC (PDF/Word/PowerPoint/text via parsers,
# images via Apple Vision OCR). Their text is sent to the model with every
# question, placed first in the prompt so OpenAI's prompt caching makes the
# repeated part cheap. If the library is bigger than the budget, only the
# excerpts most related to the on-screen question are sent.
REF_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".heic", ".heif", ".tif", ".tiff", ".gif", ".bmp", ".webp"}
REF_DOC_EXTS = {".pdf", ".docx", ".pptx"}
REF_TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".json", ".rtf", ".html", ".htm", ".py", ".tex"}
REF_FILE_FILTER = (
    "Reference files (*.pdf *.docx *.pptx *.txt *.md *.markdown *.csv *.json *.html *.htm "
    "*.png *.jpg *.jpeg *.heic *.heif *.tif *.tiff *.gif *.bmp *.webp);;All files (*)"
)


def data_dir() -> Path:
    base = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.AppDataLocation)
    if not base:
        base = str(Path.home() / ".local" / "share")
    d = Path(base)
    if d.name != APP_NAME:
        d = d / APP_NAME
    return d


def ref_kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in REF_IMAGE_EXTS:
        return "image"
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "word"
    if ext == ".pptx":
        return "slides"
    return "text"


def _tidy_text(t: str) -> str:
    t = t.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def extract_text(path: Path) -> str:
    """Return the plain text of a reference file (may be empty for scanned PDFs)."""
    kind = ref_kind(path)
    if kind == "image":
        return _tidy_text(ocr_image_file(str(path)))
    if kind == "pdf":
        from pypdf import PdfReader
        parts = []
        for i, page in enumerate(PdfReader(str(path)).pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                parts.append(f"[page {i}]\n{t}")
        return _tidy_text("\n\n".join(parts))
    if kind == "word":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return _tidy_text("\n".join(parts))
    if kind == "slides":
        from pptx import Presentation
        parts = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
            texts = [t for t in texts if t.strip()]
            if texts:
                parts.append(f"[slide {i}]\n" + "\n".join(texts))
        return _tidy_text("\n\n".join(parts))
    raw = path.read_text(encoding="utf-8", errors="replace")
    if kind == "text" and path.suffix.lower() in (".html", ".htm"):
        raw = re.sub(r"<script.*?</script>|<style.*?</style>", " ", raw, flags=re.S | re.I)
        raw = re.sub(r"<[^>]+>", " ", raw)
    return _tidy_text(raw)


_STOPWORDS = set("""a an and are as at be by for from has have in is it its of on or that the this to was
were what when where which who why will with how does do did can could should would about into than then
there these those your you not but if all any each other some such""".split())


def _terms(text: str) -> set:
    return {w for w in re.findall(r"[a-z0-9][a-z0-9'\-]{2,}", (text or "").lower()) if w not in _STOPWORDS}


def _chunks(text: str, size: int) -> List[str]:
    out, cur = [], ""
    for para in text.split("\n\n"):
        if len(cur) + len(para) + 2 > size and cur:
            out.append(cur)
            cur = para
        else:
            cur = f"{cur}\n\n{para}" if cur else para
        while len(cur) > size * 2:  # very long paragraph: hard split
            out.append(cur[:size]); cur = cur[size:]
    if cur:
        out.append(cur)
    return out


@dataclass
class RefDoc:
    id: str
    name: str
    path: str
    kind: str
    chars: int = 0
    status: str = "pending"   # pending | ready | empty | error
    error: str = ""
    added: str = ""
    text: str = field(default="", repr=False)

    def public(self) -> dict:
        return {k: v for k, v in self.__dict__.items() if k != "text"}


class ReferenceLibrary(QObject):
    """User-added reference files, their extracted text, and prompt context selection."""
    changed = pyqtSignal()
    MAX_CHARS = int(os.getenv("MM_REF_MAX_CHARS", "60000"))
    CHUNK = 1200

    def __init__(self):
        super().__init__()
        self._docs: List[RefDoc] = []
        self._lock = threading.Lock()      # guards the doc list
        self._io_lock = threading.Lock()   # serializes index writes (indexing threads save too)
        self.enabled = True
        self._load()

    # ----- persistence -----
    def _index_path(self) -> Path:
        return data_dir() / "references.json"

    def _text_path(self, doc_id: str) -> Path:
        return data_dir() / "references" / f"{doc_id}.txt"

    def _load(self):
        try:
            data = json.loads(self._index_path().read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError, OSError):
            return
        self.enabled = bool(data.get("enabled", True))
        for d in data.get("docs", []):
            try:
                doc = RefDoc(**{k: d[k] for k in ("id", "name", "path", "kind")},
                             chars=int(d.get("chars", 0)), status=d.get("status", "pending"),
                             error=d.get("error", ""), added=d.get("added", ""))
            except (KeyError, TypeError):
                continue
            tp = self._text_path(doc.id)
            if doc.status == "ready" and tp.exists():
                doc.text = tp.read_text(encoding="utf-8", errors="replace")
                doc.chars = len(doc.text)
            elif doc.status == "ready":
                doc.status = "pending"
            self._docs.append(doc)
        for doc in list(self._docs):
            if doc.status == "pending":
                self._start_index(doc)

    def _save(self):
        with self._lock:
            payload = {"enabled": self.enabled, "docs": [d.public() for d in self._docs]}
        path = self._index_path()
        with self._io_lock:
            path.parent.mkdir(parents=True, exist_ok=True)
            tmp = path.with_name(f".references.{uuid.uuid4().hex[:8]}.tmp")
            tmp.write_text(json.dumps(payload, indent=2), encoding="utf-8")
            os.replace(tmp, path)

    # ----- public API -----
    @property
    def docs(self) -> List[RefDoc]:
        with self._lock:
            return list(self._docs)

    def ready_docs(self) -> List[RefDoc]:
        return [d for d in self.docs if d.status == "ready"]

    def total_chars(self) -> int:
        return sum(d.chars for d in self.ready_docs())

    def fits_in_one_prompt(self) -> bool:
        return self.total_chars() <= self.MAX_CHARS

    def set_enabled(self, on: bool):
        self.enabled = bool(on)
        self._save()
        self.changed.emit()

    def add_paths(self, paths: List[str]) -> int:
        added = 0
        for raw in paths:
            path = Path(raw).expanduser()
            if not path.is_file():
                continue
            with self._lock:
                if any(d.path == str(path) for d in self._docs):
                    continue
                doc = RefDoc(id=uuid.uuid4().hex[:12], name=path.name, path=str(path),
                             kind=ref_kind(path), added=time.strftime("%Y-%m-%d %H:%M"))
                self._docs.append(doc)
            added += 1
            self._start_index(doc)
        if added:
            self._save()
            self.changed.emit()
        return added

    def remove(self, doc_id: str):
        with self._lock:
            self._docs = [d for d in self._docs if d.id != doc_id]
        try:
            self._text_path(doc_id).unlink()
        except OSError:
            pass
        self._save()
        self.changed.emit()

    def clear(self):
        with self._lock:
            ids = [d.id for d in self._docs]
            self._docs = []
        for i in ids:
            try:
                self._text_path(i).unlink()
            except OSError:
                pass
        self._save()
        self.changed.emit()

    # ----- indexing (background) -----
    def _start_index(self, doc: RefDoc):
        threading.Thread(target=self._index, args=(doc,), daemon=True).start()

    def _index(self, doc: RefDoc):
        try:
            text = extract_text(Path(doc.path))
            with self._lock:
                doc.text = text
                doc.chars = len(text)
                doc.status = "ready" if len(text) >= 20 else "empty"
                doc.error = ""
            if doc.status == "ready":
                tp = self._text_path(doc.id)
                tp.parent.mkdir(parents=True, exist_ok=True)
                tp.write_text(text, encoding="utf-8")
        except Exception as e:
            with self._lock:
                doc.status = "error"
                doc.error = str(e)[:200]
        self._save()
        self.changed.emit()

    # ----- prompt context -----
    def context_for(self, query: str) -> Tuple[str, List[str]]:
        """Return (reference block, file names used) for a question. Empty if disabled/none."""
        docs = self.ready_docs()
        if not self.enabled or not docs:
            return "", []
        if sum(d.chars for d in docs) <= self.MAX_CHARS:
            block = "\n\n".join(f"### {d.name}\n{d.text}" for d in docs)
            return block, [d.name for d in docs]
        # too big for one prompt: pick the chunks that overlap the question most
        q = _terms(query)
        scored = []
        for d in docs:
            for ch in _chunks(d.text, self.CHUNK):
                ct = _terms(ch)
                sc = sum(2 if len(t) >= 6 else 1 for t in (q & ct))
                scored.append((sc, d.name, ch))
        scored.sort(key=lambda x: -x[0])
        picked, used, names = [], 0, []
        for sc, name, ch in scored:
            if sc <= 0 and picked:
                break
            if used + len(ch) > self.MAX_CHARS:
                continue
            picked.append((name, ch)); used += len(ch)
            if name not in names:
                names.append(name)
        block = "\n\n".join(f"### {name} (excerpt)\n{ch}" for name, ch in picked)
        return block, names


_refs: Optional[ReferenceLibrary] = None


def get_refs() -> ReferenceLibrary:
    global _refs
    if _refs is None:
        _refs = ReferenceLibrary()
    return _refs


# ================== theme + liquid glass primitives ==================
DEFAULT_PANE_COLOR = QColor(212, 175, 55, 230)
PANE_COLOR = QColor(DEFAULT_PANE_COLOR)
SHOW_DEBUG_STATUS = False


def is_dark() -> bool:
    try:
        return QGuiApplication.styleHints().colorScheme() == Qt.ColorScheme.Dark
    except Exception:
        return QGuiApplication.palette().window().color().lightness() < 128


@dataclass
class Palette:
    dark: bool
    fill: QColor          # emulated glass fill (no native blur)
    fill_blur: QColor     # light tint layered over native blur
    border: QColor        # hairline edge
    highlight: QColor     # specular top-left stroke
    text: str
    subtext: str
    accent: str = "#0a84ff"
    ok: str = "#1aa34a"
    warn: str = "#e0a800"
    danger: str = "#ff453a"
    neutral: str = "#8e8e93"
    button_bg: str = "rgba(255,255,255,0.38)"
    button_bg_hover: str = "rgba(255,255,255,0.55)"
    button_border: str = "rgba(0,0,0,0.10)"
    button_border_hover: str = "rgba(0,0,0,0.16)"
    btn_top: str = "rgba(255,255,255,0.88)"      # capsule sheen: bright rim at the top …
    btn_mid: str = "rgba(255,255,255,0.46)"
    btn_bottom: str = "rgba(255,255,255,0.34)"   # … settling into the frosted body
    btn_top_hover: str = "rgba(255,255,255,0.98)"
    btn_bottom_hover: str = "rgba(255,255,255,0.55)"
    btn_pressed: str = "rgba(0,0,0,0.07)"
    field_bg: str = "rgba(255,255,255,0.45)"
    row_bg: str = "rgba(255,255,255,0.28)"
    row_bg_hover: str = "rgba(255,255,255,0.45)"
    badge_bg: str = "rgba(120,120,128,0.16)"


def current_palette() -> Palette:
    if is_dark():
        return Palette(
            dark=True,
            fill=QColor(30, 30, 36, 200),
            fill_blur=QColor(30, 30, 36, 90),
            border=QColor(255, 255, 255, 48),
            highlight=QColor(255, 255, 255, 120),
            text="#f5f5f7",
            subtext="rgba(235,235,245,0.62)",
            ok="#30d158",
            warn="#ffd60a",
            button_bg="rgba(255,255,255,0.10)",
            button_bg_hover="rgba(255,255,255,0.18)",
            button_border="rgba(255,255,255,0.16)",
            button_border_hover="rgba(255,255,255,0.28)",
            btn_top="rgba(255,255,255,0.28)",
            btn_mid="rgba(255,255,255,0.13)",
            btn_bottom="rgba(255,255,255,0.08)",
            btn_top_hover="rgba(255,255,255,0.36)",
            btn_bottom_hover="rgba(255,255,255,0.16)",
            btn_pressed="rgba(255,255,255,0.05)",
            field_bg="rgba(255,255,255,0.08)",
            row_bg="rgba(255,255,255,0.06)",
            row_bg_hover="rgba(255,255,255,0.12)",
            badge_bg="rgba(255,255,255,0.12)",
        )
    return Palette(
        dark=False,
        fill=QColor(255, 255, 255, 200),
        fill_blur=QColor(255, 255, 255, 95),
        border=QColor(0, 0, 0, 34),
        highlight=QColor(255, 255, 255, 230),
        text="#1d1d1f",
        subtext="rgba(60,60,67,0.68)",
    )


def _level_qss(pal: "Palette") -> str:
    """Color variants keyed by the dynamic `level` property (see confidence_level)."""
    colors = {"high": pal.ok, "risky": pal.warn, "low": pal.danger}
    out = []
    for level, color in colors.items():
        for name in ("headline", "headline_md", "headline_small", "conftext"):
            out.append(f'QLabel#{name}[level="{level}"] {{ color: {color}; }}')
        out.append(f'QLabel#confdot[level="{level}"] {{ background: {color}; }}')
    return "\n    ".join(out)


def glass_qss(pal: Palette) -> str:
    return f"""
    QWidget {{
        background: transparent;
        color: {pal.text};
        font-size: 13px;
    }}
    QLabel#title {{ font-size: 15px; font-weight: 600; }}
    QLabel#subtext {{ color: {pal.subtext}; font-size: 12.5px; }}
    QLabel#headline {{ color: {pal.text}; font-size: 21px; font-weight: 700; }}
    QLabel#headline_md {{ color: {pal.text}; font-size: 17px; font-weight: 700; }}
    QLabel#headline_small {{ color: {pal.text}; font-size: 14px; font-weight: 600; }}
    QLabel#error {{ color: {pal.danger}; font-size: 12px; }}
    QLabel#confdot {{ background: {pal.neutral}; border-radius: 5px; min-width: 10px; max-width: 10px; min-height: 10px; max-height: 10px; }}
    QLabel#conftext {{ color: {pal.neutral}; font-size: 12px; font-weight: 600; }}
    {_level_qss(pal)}
    QLabel#badge {{
        background: {pal.badge_bg}; color: {pal.subtext};
        border-radius: 9px; padding: 1px 8px; font-size: 11.5px;
    }}
    QLabel#time {{ color: {pal.subtext}; font-size: 11.5px; }}
    QLabel#dot {{ border-radius: 5px; }}
    /* capsule buttons: full-round ends, a bright top rim fading into frosted glass */
    QPushButton {{
        background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
            stop:0 {pal.btn_top}, stop:0.5 {pal.btn_mid}, stop:1 {pal.btn_bottom});
        border: 1px solid {pal.button_border};
        border-radius: 17px;
        min-height: 34px;
        padding: 0 18px;
        font-weight: 600;
    }}
    QPushButton:hover {{
        background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
            stop:0 {pal.btn_top_hover}, stop:0.5 {pal.btn_mid}, stop:1 {pal.btn_bottom_hover});
        border: 1px solid {pal.button_border_hover};
    }}
    QPushButton:pressed {{ background: {pal.btn_pressed}; }}
    QPushButton:disabled {{ color: {pal.subtext}; }}
    QPushButton#primary {{
        background: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 #62adff, stop:1 {pal.accent});
        color: white; border: 1px solid rgba(255,255,255,0.40);
    }}
    QPushButton#primary:hover {{
        background: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 #7cbbff, stop:1 #2b93ff);
    }}
    QPushButton#primary:pressed {{ background: #0a6fd6; }}
    QPushButton#danger {{
        background: qlineargradient(x1:0, y1:0, x2:0, y2:1, stop:0 rgba(255,110,100,0.34), stop:1 rgba(255,69,58,0.16));
        color: {pal.danger}; border: 1px solid rgba(255,69,58,0.32);
    }}
    QPushButton#danger:hover {{ background: rgba(255,69,58,0.32); }}
    QPushButton#ghost {{
        background: transparent; border: none; min-height: 0; padding: 2px 6px;
        border-radius: 11px; color: {pal.subtext}; font-size: 14px;
    }}
    QPushButton#ghost:hover {{ background: {pal.button_bg}; }}
    QPushButton#small {{ min-height: 26px; border-radius: 13px; padding: 0 12px; font-size: 12px; }}
    QLineEdit {{
        background: {pal.field_bg};
        border: 1px solid {pal.button_border};
        border-radius: 17px;
        min-height: 34px;
        padding: 0 14px;
        selection-background-color: {pal.accent};
    }}
    QLineEdit:focus {{ border: 1px solid {pal.accent}; }}
    QCheckBox {{ spacing: 8px; }}
    QCheckBox::indicator {{ width: 18px; height: 18px; border-radius: 9px;
        border: 1px solid {pal.button_border}; background: {pal.field_bg}; }}
    QCheckBox::indicator:checked {{ background: {pal.accent}; border: 1px solid {pal.accent}; }}
    QScrollArea {{ border: none; background: transparent; }}
    QScrollBar:vertical {{ background: transparent; width: 8px; margin: 2px; }}
    QScrollBar::handle:vertical {{ background: {pal.badge_bg}; border-radius: 4px; min-height: 24px; }}
    QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{ height: 0; }}
    QFrame#row {{ background: {pal.row_bg}; border-radius: 10px; }}
    QFrame#row:hover {{ background: {pal.row_bg_hover}; }}
    """


class GlassCard(QFrame):
    """Rounded translucent panel: frosted fill, liquid gradient, hairline edge, specular highlight."""

    def __init__(self, parent=None, radius: int = 24):
        super().__init__(parent)
        self.radius = radius
        self._blur_active = False
        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, False)

    def set_blur_active(self, on: bool):
        self._blur_active = on
        self.update()

    def paintEvent(self, _):
        pal = current_palette()
        p = QPainter(self)
        p.setRenderHint(QPainter.RenderHint.Antialiasing, True)
        r = QRectF(self.rect()).adjusted(0.5, 0.5, -0.5, -0.5)
        path = QPainterPath()
        path.addRoundedRect(r, self.radius, self.radius)

        # frosted fill (lighter when the OS is blurring behind us)
        p.fillPath(path, pal.fill_blur if self._blur_active else pal.fill)

        # "liquid" sheen: bright at the top, clear in the middle, faint at the bottom
        g = QLinearGradient(r.topLeft(), r.bottomLeft())
        top_a, bot_a = (26, 6) if pal.dark else (80, 22)
        g.setColorAt(0.0, QColor(255, 255, 255, top_a))
        g.setColorAt(0.45, QColor(255, 255, 255, 0))
        g.setColorAt(1.0, QColor(255, 255, 255, bot_a))
        p.fillPath(path, g)

        # hairline edge
        p.setPen(QPen(pal.border, 1))
        p.setBrush(Qt.BrushStyle.NoBrush)
        p.drawPath(path)

        # specular highlight along the top-left, fading out
        inner = QPainterPath()
        inner.addRoundedRect(r.adjusted(1, 1, -1, -1), self.radius - 1, self.radius - 1)
        hg = QLinearGradient(r.topLeft(), r.bottomRight())
        hg.setColorAt(0.0, pal.highlight)
        hg.setColorAt(0.55, QColor(255, 255, 255, 0))
        p.setPen(QPen(QBrush(hg), 1.2))
        p.drawPath(inner)


# ---------- native macOS blur (NSVisualEffectView behind Qt's content) ----------
try:
    import objc
    from AppKit import NSVisualEffectView, NSColor
    from Foundation import NSMakeRect
    _HAVE_OBJC = True
except Exception:
    _HAVE_OBJC = False

_NS_MATERIAL_POPOVER = 6
_NS_BLEND_BEHIND_WINDOW = 0
_NS_STATE_ACTIVE = 1
_NS_WINDOW_BELOW = -1
_NS_VIEW_WIDTH_SIZABLE = 2
_NS_VIEW_HEIGHT_SIZABLE = 16


def _appkit_rect(view, rect: QRect):
    """Qt top-left window coords -> AppKit bottom-left coords inside the same view."""
    h = view.frame().size.height
    return NSMakeRect(rect.x(), h - (rect.y() + rect.height()), rect.width(), rect.height())


def apply_native_blur(widget: QWidget, radius: int, rect: Optional[QRect] = None):
    """Slide an NSVisualEffectView underneath the Qt view of a translucent window.
    Returns the effect view (keep it to re-frame later) or None if unavailable."""
    # winId() is only an NSView* on the real Cocoa platform (not offscreen/minimal)
    if not _HAVE_OBJC or sys.platform != "darwin" or QGuiApplication.platformName() != "cocoa":
        return None
    try:
        view = objc.objc_object(c_void_p=int(widget.winId()))
        win = view.window()
        if win is None:
            return None
        frame_view = view.superview()
        if frame_view is None:
            return None
        effect = NSVisualEffectView.alloc().initWithFrame_(view.frame())
        effect.setMaterial_(_NS_MATERIAL_POPOVER)
        effect.setBlendingMode_(_NS_BLEND_BEHIND_WINDOW)
        effect.setState_(_NS_STATE_ACTIVE)
        effect.setWantsLayer_(True)
        effect.layer().setCornerRadius_(float(radius))
        effect.layer().setMasksToBounds_(True)
        if rect is None:
            effect.setAutoresizingMask_(_NS_VIEW_WIDTH_SIZABLE | _NS_VIEW_HEIGHT_SIZABLE)
        else:
            effect.setFrame_(_appkit_rect(view, rect))
        frame_view.addSubview_positioned_relativeTo_(effect, _NS_WINDOW_BELOW, view)
        win.setOpaque_(False)
        win.setBackgroundColor_(NSColor.clearColor())
        return effect
    except Exception as e:
        print("[MirrorMirror] native blur unavailable:", e)
        return None


def reframe_native_blur(widget: QWidget, effect, rect: Optional[QRect], visible: bool = True):
    if effect is None:
        return
    try:
        effect.setHidden_(not visible)
        if rect is not None and visible:
            view = objc.objc_object(c_void_p=int(widget.winId()))
            effect.setFrame_(_appkit_rect(view, rect))
    except Exception:
        pass


class _GlassChrome:
    """Mixin: frameless translucent top-level made of one GlassCard with a title row.
    Drag anywhere on the card to move. Native blur attached on first show."""
    RADIUS = 24

    def _init_glass(self, title: str, closable: bool = True, min_width: int = 0):
        self.setWindowFlags(self.windowFlags() | Qt.WindowType.FramelessWindowHint)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground, True)
        self.setStyleSheet(glass_qss(current_palette()))
        if min_width:
            self.setMinimumWidth(min_width)

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        self.card = GlassCard(self, radius=self.RADIUS)
        root.addWidget(self.card)

        cl = QVBoxLayout(self.card)
        cl.setContentsMargins(20, 14, 20, 18)
        cl.setSpacing(12)

        self.title_row = QHBoxLayout()
        self.title_row.setSpacing(8)
        self.title_lbl = QLabel(title)
        self.title_lbl.setObjectName("title")
        self.title_row.addWidget(self.title_lbl)
        self.title_row.addStretch(1)
        self.close_btn = None
        if closable:
            self.close_btn = QPushButton("✕")
            self.close_btn.setObjectName("ghost")
            self.close_btn.setFixedSize(26, 26)
            self.close_btn.setCursor(Qt.CursorShape.PointingHandCursor)
            self.close_btn.clicked.connect(self.close)
            self.title_row.addWidget(self.close_btn)
        cl.addLayout(self.title_row)

        self.body = QVBoxLayout()
        self.body.setSpacing(10)
        cl.addLayout(self.body)

        self._drag_offset = None
        self._blur = None

    def showEvent(self, e):
        super().showEvent(e)
        if self._blur is None:
            QTimer.singleShot(0, self._attach_blur)

    def _attach_blur(self):
        if self._blur is None:
            self._blur = apply_native_blur(self, self.RADIUS)
            self.card.set_blur_active(self._blur is not None)

    def mousePressEvent(self, e):
        if e.button() == Qt.MouseButton.LeftButton:
            self._drag_offset = e.globalPosition().toPoint() - self.frameGeometry().topLeft()
            e.accept()

    def mouseMoveEvent(self, e):
        if self._drag_offset is not None and (e.buttons() & Qt.MouseButton.LeftButton):
            self.move(e.globalPosition().toPoint() - self._drag_offset)
            e.accept()

    def mouseReleaseEvent(self, e):
        self._drag_offset = None


class GlassWindow(_GlassChrome, QWidget):
    def __init__(self, title: str, closable: bool = True, min_width: int = 0):
        QWidget.__init__(self)
        self._init_glass(title, closable, min_width)


class GlassDialog(_GlassChrome, QDialog):
    def __init__(self, title: str, parent=None, min_width: int = 0):
        QDialog.__init__(self, parent)
        self.setModal(True)
        self._init_glass(title, closable=True, min_width=min_width)


# ================== API key dialog ==================
class ApiKeyDialog(GlassDialog):
    """Asks the user for an OpenAI API key and stores it in the user config file."""

    def __init__(self, parent=None, current_key: str = ""):
        super().__init__("OpenAI API Key", parent, min_width=480)

        intro = QLabel(
            "Mirror Mirror needs an OpenAI API key to answer questions. "
            "Create one at platform.openai.com → API keys, then paste it below."
        )
        intro.setWordWrap(True)
        self.body.addWidget(intro)

        where = QLabel(f"Stored only on this Mac:\n{config_path()}")
        where.setObjectName("subtext")
        where.setWordWrap(True)
        self.body.addWidget(where)

        self._edit = QLineEdit()
        self._edit.setPlaceholderText("sk-…")
        self._edit.setEchoMode(QLineEdit.EchoMode.Password)
        self._edit.setText(current_key)
        self.body.addWidget(self._edit)

        self._show_chk = QCheckBox("Show key")
        self._show_chk.toggled.connect(
            lambda on: self._edit.setEchoMode(
                QLineEdit.EchoMode.Normal if on else QLineEdit.EchoMode.Password
            )
        )
        self.body.addWidget(self._show_chk)

        self._error = QLabel("")
        self._error.setObjectName("error")
        self._error.setWordWrap(True)
        self.body.addWidget(self._error)

        row = QHBoxLayout()
        row.addStretch(1)
        cancel = QPushButton("Cancel")
        cancel.clicked.connect(self.reject)
        self._save_btn = QPushButton("Save")
        self._save_btn.setObjectName("primary")
        self._save_btn.setDefault(True)
        self._save_btn.clicked.connect(self._on_save)
        row.addWidget(cancel)
        row.addWidget(self._save_btn)
        self.body.addLayout(row)

        self._edit.returnPressed.connect(self._on_save)
        self._edit.setFocus()

    def key(self) -> str:
        return self._edit.text().strip()

    def _on_save(self):
        key = self.key()
        if not key:
            self._error.setText("Please paste an API key.")
            return
        if any(ch.isspace() for ch in key):
            self._error.setText("The key should not contain spaces or line breaks.")
            return

        self._error.setText("Checking key with OpenAI…")
        self._save_btn.setEnabled(False)
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        QApplication.processEvents()
        try:
            err = verify_api_key(key)
        finally:
            QApplication.restoreOverrideCursor()
            self._save_btn.setEnabled(True)

        if err:
            self._error.setText(err)
            return

        set_api_key(key)
        self.accept()


def ensure_api_key(parent=None) -> bool:
    """Prompt for a key if none is configured. Returns True when a key is available."""
    if get_api_key():
        return True
    dlg = ApiKeyDialog(parent)
    return dlg.exec() == QDialog.DialogCode.Accepted


# ================== helpers ==================
@dataclass
class CaptureRect:
    left:int; top:int; width:int; height:int

def pil_from_qimage(qimg: QImage) -> Image.Image:
    qimg = qimg.convertToFormat(QImage.Format.Format_RGBA8888)
    width = qimg.width()
    height = qimg.height()
    ptr = qimg.bits()
    ptr.setsize(qimg.sizeInBytes())
    arr = bytes(ptr)
    return Image.frombuffer("RGBA", (width, height), arr, "raw", "RGBA", 0, 1).convert("RGB")


def mask_image(qimg: QImage, rects: List[QRect], pad: int = 3) -> None:
    """Blank the given pane-local logical rects out of a captured image (in place).
    This is what keeps the on-screen answer card from ever reaching OCR."""
    if not rects:
        return
    # QPainter works in logical coords and applies the image's devicePixelRatio
    # itself (Retina captures are 2x), so the rects are used as-is.
    p = QPainter(qimg)
    try:
        p.setCompositionMode(QPainter.CompositionMode.CompositionMode_Source)
        for r in rects:
            p.fillRect(r.adjusted(-pad, -pad, pad, pad), QColor(255, 255, 255, 255))
    finally:
        p.end()


# ================== answers ==================
@dataclass
class Answer:
    headline: str
    confidence: str            # "92%" or ""
    detail: str
    raw: str
    conf_value: Optional[int] = None
    source: str = ""             # "notes.pdf" / "general knowledge" / ""
    when: str = field(default_factory=lambda: time.strftime("%H:%M"))


_CONF_RE = re.compile(r"confidence[^0-9]{0,12}(\d{1,3})\s*%", re.I)

# Confidence tiers (percent). Green = high, yellow = risky, red = needs more reference.
CONF_HIGH = int(os.getenv("MM_CONF_HIGH", "90"))
CONF_RISKY = int(os.getenv("MM_CONF_RISKY", "70"))
CONF_TOOLTIPS = {
    "high": f"High confidence (≥ {CONF_HIGH}%)",
    "risky": f"Risky ({CONF_RISKY}–{CONF_HIGH - 1}%)",
    "low": f"Needs more reference (< {CONF_RISKY}%)",
    "unknown": "No confidence given",
}


def confidence_level(value: Optional[int]) -> str:
    if value is None:
        return "unknown"
    if value >= CONF_HIGH:
        return "high"
    if value >= CONF_RISKY:
        return "risky"
    return "low"


def set_level(widget: QWidget, level: str) -> None:
    """Set the dynamic `level` property and re-polish so the stylesheet variant applies."""
    if widget.property("level") == level:
        return
    widget.setProperty("level", level)
    st = widget.style()
    st.unpolish(widget)
    st.polish(widget)
    widget.update()


def parse_answer(raw: str) -> Answer:
    text = (raw or "").strip()
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    m = _CONF_RE.search(text)
    value = min(100, int(m.group(1))) if m else None
    conf = f"{value}%" if value is not None else ""
    headline = lines[0] if lines else ""
    rest, source = [], ""
    for l in lines[1:]:
        if _CONF_RE.search(l) and len(l) <= 40:
            continue  # the "Confidence: NN%" line itself
        ms = re.match(r"^\**\s*source\s*:\s*(.+?)\s*\**$", l, re.I)
        if ms and len(l) <= 120:
            source = ms.group(1).strip()
            continue
        rest.append(l)
    # strip any inline confidence + markdown noise from the headline
    headline = _CONF_RE.sub("", headline)
    headline = re.sub(r"\(\s*\)|\[\s*\]", "", headline)   # empty brackets left behind
    headline = re.sub(r"[\s(\[]+$", "", headline)           # dangling opener at the end
    headline = headline.strip(" *#-–:_\t")
    if len(headline) > 80:
        headline = headline[:77].rstrip() + "…"
    return Answer(headline=headline or "(no answer)", confidence=conf,
                  detail="\n".join(rest), raw=text, conf_value=value, source=source)


def _is_error_text(s: str) -> bool:
    s = (s or "").strip()
    return s.startswith("(") and "error" in s.lower()


class ConfidenceBadge(QWidget):
    """Colored dot + optional percentage. Green/yellow/red by confidence_level."""

    def __init__(self, parent=None, show_text: bool = True):
        super().__init__(parent)
        self._show_text = show_text
        self._level = "unknown"
        self._text = ""
        h = QHBoxLayout(self)
        h.setContentsMargins(0, 0, 0, 0)
        h.setSpacing(5)
        self._dot = QLabel(); self._dot.setObjectName("confdot")
        self._dot.setFixedSize(10, 10)
        self._txt = QLabel(""); self._txt.setObjectName("conftext")
        h.addWidget(self._dot, 0, Qt.AlignmentFlag.AlignVCenter)
        h.addWidget(self._txt, 0, Qt.AlignmentFlag.AlignVCenter)
        self._apply()

    @property
    def level(self) -> str:
        return self._level

    def set_answer(self, ans: Optional[Answer]):
        if ans is None:
            self.set_level("unknown", "")
        else:
            self.set_level(confidence_level(ans.conf_value), ans.confidence)

    def set_level(self, level: str, text: str = ""):
        self._level = level
        self._text = text
        self._apply()

    def set_show_text(self, on: bool):
        if on != self._show_text:
            self._show_text = on
            self._apply()

    def _apply(self):
        set_level(self._dot, self._level)
        set_level(self._txt, self._level)
        self._txt.setText(self._text)
        self._txt.setVisible(self._show_text and bool(self._text))
        self.setToolTip(CONF_TOOLTIPS.get(self._level, ""))


class AnswerRow(QFrame):
    """One compact history entry; click to reveal the full text."""

    def __init__(self, ans: Answer, parent=None):
        super().__init__(parent)
        self.setObjectName("row")
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self._ans = ans
        v = QVBoxLayout(self)
        v.setContentsMargins(10, 6, 10, 6)
        v.setSpacing(4)
        top = QHBoxLayout()
        top.setSpacing(8)
        t = QLabel(ans.when); t.setObjectName("time")
        self._head = QLabel(ans.headline); self._head.setObjectName("headline_small")
        self._head.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        set_level(self._head, confidence_level(ans.conf_value))
        self._badge = ConfidenceBadge(show_text=True)
        self._badge.set_answer(ans)
        top.addWidget(t)
        top.addWidget(self._head, 1)
        top.addWidget(self._badge)
        v.addLayout(top)
        self._full = QLabel(ans.raw)
        self._full.setObjectName("subtext")
        self._full.setWordWrap(True)
        self._full.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        self._full.hide()
        v.addWidget(self._full)

    def mousePressEvent(self, e):
        if e.button() == Qt.MouseButton.LeftButton:
            self._full.setVisible(self._full.isHidden())
        e.accept()

    def mouseMoveEvent(self, e): e.accept()
    def mouseReleaseEvent(self, e): e.accept()


class AnswerCard(GlassCard):
    """Liquid-glass answer panel that lives INSIDE the capture pane.
    The pane reports this widget's rect so the engine blanks it before OCR.

    Responsive tiers (chosen by the pane from its own size):
      full     headline · dot+% · explanation · toolbar · history
      compact  headline (smaller) · dot+% · toolbar
      mini     one-line pill: dot · headline · count   (percentage/count drop when narrow)
      hidden   card hidden; the toolbar window shows the latest answer instead
    """
    MAX_HISTORY = 25
    PILL_H = 34
    COMPACT_H = 84

    layout_changed = pyqtSignal()
    dock_toggle_requested = pyqtSignal()

    @staticmethod
    def tier_for(pane_w: int, pane_h: int) -> str:
        if pane_h < 120 or pane_w < 200:
            return "hidden"
        if pane_h >= 320 and pane_w >= 420:
            return "full"
        if pane_h >= 200 and pane_w >= 300:
            return "compact"
        return "mini"

    def __init__(self, parent=None):
        super().__init__(parent, radius=18)
        self.setStyleSheet(glass_qss(current_palette()))
        self.setCursor(Qt.CursorShape.ArrowCursor)
        self._answers: List[Answer] = []
        self._collapsed = False      # the user's choice
        self._details = False        # explanation + history only when asked, so the
                                     # card stays small and covers as little as possible
        self._tier = "full"
        self._pane_w = 10_000

        outer = QVBoxLayout(self)
        outer.setContentsMargins(14, 10, 14, 10)
        outer.setSpacing(6)

        # ----- collapsed pill -----
        self._pill = QWidget()
        pl = QHBoxLayout(self._pill)
        pl.setContentsMargins(0, 0, 0, 0)
        pl.setSpacing(8)
        self._pill_btn = QPushButton("▸"); self._pill_btn.setObjectName("ghost")
        self._pill_btn.setFixedSize(22, 22); self._pill_btn.setToolTip("Expand")
        self._pill_btn.clicked.connect(self.toggle)
        self._pill_conf = ConfidenceBadge(show_text=True)
        self._pill_head = QLabel(""); self._pill_head.setObjectName("headline_small")
        self._pill_head.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        self._pill_count = QLabel(""); self._pill_count.setObjectName("badge")
        pl.addWidget(self._pill_btn)
        pl.addWidget(self._pill_conf)
        pl.addWidget(self._pill_head, 1)
        pl.addWidget(self._pill_count)
        outer.addWidget(self._pill)
        self._pill.hide()

        # ----- expanded content (full / compact) -----
        self._full = QWidget()
        fl = QVBoxLayout(self._full)
        fl.setContentsMargins(0, 0, 0, 0)
        fl.setSpacing(6)

        head_row = QHBoxLayout(); head_row.setSpacing(8)
        self._headline = QLabel("Waiting for text…"); self._headline.setObjectName("headline")
        self._headline.setWordWrap(True)
        self._headline.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        self._headline.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        self._conf = ConfidenceBadge(show_text=True); self._conf.hide()
        head_row.addWidget(self._headline, 1)
        head_row.addWidget(self._conf, 0, Qt.AlignmentFlag.AlignTop)
        fl.addLayout(head_row)

        self._detail = QLabel(""); self._detail.setObjectName("subtext")
        self._detail.setWordWrap(True)
        self._detail.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        self._detail.hide()
        fl.addWidget(self._detail)

        self._source = QLabel(""); self._source.setObjectName("subtext")
        self._source.hide()
        fl.addWidget(self._source)

        self._error = QLabel(""); self._error.setObjectName("error")
        self._error.setWordWrap(True); self._error.hide()
        fl.addWidget(self._error)

        tools = QHBoxLayout(); tools.setSpacing(6)
        self._chev = QPushButton("▾"); self._chev.setObjectName("ghost")
        self._chev.setFixedSize(22, 22); self._chev.setToolTip("Collapse")
        self._chev.clicked.connect(self.toggle)
        self._count = QLabel("0 answers"); self._count.setObjectName("badge")
        self._details_btn = QPushButton("Details"); self._details_btn.setObjectName("small")
        self._details_btn.setToolTip("Show the explanation and earlier answers")
        self._details_btn.clicked.connect(self.toggle_details)
        self._copy_btn = QPushButton("Copy"); self._copy_btn.setObjectName("small")
        self._copy_btn.clicked.connect(self.copy_latest)
        self._clear_btn = QPushButton("Clear"); self._clear_btn.setObjectName("small")
        self._clear_btn.clicked.connect(self.clear)
        self._dock_btn = QPushButton("⇅"); self._dock_btn.setObjectName("ghost")
        self._dock_btn.setFixedSize(22, 22); self._dock_btn.setToolTip("Move the card to the other edge")
        self._dock_btn.clicked.connect(self.dock_toggle_requested.emit)
        tools.addWidget(self._chev)
        tools.addWidget(self._count)
        tools.addStretch(1)
        tools.addWidget(self._details_btn)
        tools.addWidget(self._copy_btn)
        tools.addWidget(self._clear_btn)
        tools.addWidget(self._dock_btn)
        fl.addLayout(tools)

        self._scroll = QScrollArea()
        self._scroll.setWidgetResizable(True)
        self._scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self._scroll.setFrameShape(QFrame.Shape.NoFrame)
        self._hist_host = QWidget()
        self._hist = QVBoxLayout(self._hist_host)
        self._hist.setContentsMargins(0, 0, 4, 0)
        self._hist.setSpacing(4)
        self._hist.addStretch(1)
        self._scroll.setWidget(self._hist_host)
        self._scroll.setMaximumHeight(150)
        self._scroll.hide()
        fl.addWidget(self._scroll)

        outer.addWidget(self._full)
        self._apply_state()

    # ----- state -----
    @property
    def answers(self) -> List[Answer]:
        return list(self._answers)

    @property
    def latest(self) -> Optional[Answer]:
        return self._answers[0] if self._answers else None

    @property
    def tier(self) -> str:
        return self._tier

    @property
    def collapsed(self) -> bool:
        """Effective state: the user's choice, forced to the pill in the mini tier."""
        return self._tier == "mini" or self._collapsed

    @property
    def level(self) -> str:
        return confidence_level(self.latest.conf_value) if self.latest else "unknown"

    def push(self, raw: str) -> bool:
        """Add a new answer. Returns False if ignored (empty or duplicate of latest)."""
        text = (raw or "").strip()
        if not text:
            return False
        if self._answers and self._answers[0].raw == text:
            return False
        ans = parse_answer(text)
        self._answers.insert(0, ans)
        self._hist.insertWidget(0, AnswerRow(ans))
        while len(self._answers) > self.MAX_HISTORY:
            self._answers.pop()
            item = self._hist.takeAt(self._hist.count() - 2)  # keep trailing stretch
            if item and item.widget():
                item.widget().deleteLater()
        self._error.hide()
        self._refresh()
        return True

    def set_error(self, msg: str):
        self._error.setText(msg)
        self._error.show()
        self.layout_changed.emit()

    def clear(self):
        self._answers.clear()
        while self._hist.count() > 1:
            item = self._hist.takeAt(0)
            if item and item.widget():
                item.widget().deleteLater()
        self._error.hide()
        self._refresh()

    def copy_latest(self):
        if self.latest:
            QApplication.clipboard().setText(self.latest.raw)

    def toggle(self):
        if self._tier == "mini":
            return  # nothing bigger fits
        self.set_collapsed(not self._collapsed)

    @property
    def details(self) -> bool:
        return self._details

    def toggle_details(self):
        self.set_details(not self._details)

    def set_details(self, on: bool):
        if on == self._details:
            return
        self._details = on
        self._refresh()

    def set_collapsed(self, on: bool):
        if on == self._collapsed:
            return
        self._collapsed = on
        self._refresh()

    def set_tier(self, tier: str, pane_w: int):
        """Called by the pane from its layout pass; does not emit layout_changed."""
        if tier == self._tier and pane_w == self._pane_w:
            return
        self._tier = tier
        self._pane_w = pane_w
        self._apply_state()

    def desired_height(self, pane_height: int) -> int:
        if self._tier == "hidden":
            return 0
        if self.collapsed:
            return self.PILL_H
        if self._tier == "compact":
            return self.COMPACT_H
        want = self._full.sizeHint().height() + 20
        cap = max(120, int(pane_height * 0.45))
        return max(min(want, cap), 96)

    # ----- rendering -----
    def _refresh(self):
        n = len(self._answers)
        count = f"{n} answer{'s' if n != 1 else ''}"
        self._count.setText(count)
        self._pill_count.setText(count)
        a = self.latest
        lvl = self.level
        if a:
            self._headline.setText(a.headline)
            self._detail.setText(a.detail)
            self._source.setText(f"📎 {a.source}" if a.source else "")
        else:
            self._headline.setText("Waiting for text…")
            self._detail.setText("")
            self._source.setText("")
        self._conf.set_answer(a)
        self._pill_conf.set_answer(a)
        for w in (self._headline, self._pill_head):
            set_level(w, lvl)
        self._copy_btn.setEnabled(n > 0)
        self._clear_btn.setEnabled(n > 0)
        self._apply_state()
        self.layout_changed.emit()

    def _apply_state(self):
        """Show/hide pieces for the current tier + collapsed state (no signal)."""
        mini = self._tier == "mini"
        compact = self._tier == "compact"
        collapsed = self.collapsed
        n = len(self._answers)
        self._pill.setVisible(collapsed)
        self._full.setVisible(not collapsed)
        # pill: drop the expand chevron in mini (nothing to expand), then the
        # percentage, then the count as the pane gets narrower
        self._pill_btn.setVisible(not mini)
        self._pill_conf.set_show_text(self._pane_w >= 300)
        self._pill_conf.setVisible(bool(self.latest))
        self._pill_count.setVisible(self._pane_w >= 240)
        # expanded: explanation + history only in the full tier AND only when the
        # user asked for details; otherwise the card stays one headline tall
        show_details = (not compact) and self._details
        self._conf.setVisible(bool(self.latest))
        self._detail.setVisible(show_details and bool(self._detail.text()))
        self._source.setVisible((not compact) and bool(self._source.text()))
        self._scroll.setVisible(show_details and n > 0)
        self._details_btn.setVisible(not compact)
        self._details_btn.setText("Hide" if self._details else "Details")
        self._details_btn.setEnabled(n > 0)
        want_name = "headline_md" if compact else "headline"
        if self._headline.objectName() != want_name:
            self._headline.setObjectName(want_name)
            st = self._headline.style(); st.unpolish(self._headline); st.polish(self._headline)
        self._elide_pill()

    def _elide_pill(self):
        text = self.latest.headline if self.latest else "Waiting for text…"
        fm = QFontMetrics(self._pill_head.font())
        w = max(40, self._pill_head.width() - 4)
        self._pill_head.setText(fm.elidedText(text, Qt.TextElideMode.ElideRight, w))

    def resizeEvent(self, e):
        super().resizeEvent(e)
        self._elide_pill()

    # swallow mouse so the pane's drag/click-to-OCR handlers never see card clicks
    def mousePressEvent(self, e): e.accept()
    def mouseMoveEvent(self, e): e.accept()
    def mouseReleaseEvent(self, e): e.accept()
    def mouseDoubleClickEvent(self, e):
        self.toggle(); e.accept()


# ================== transparent glass pane with perimeter drag + grips ==================
class GlassPane(QWidget):
    """Gold-framed, draggable/resizable overlay.
       - Drag from center or any border to move
       - Single-click near center to trigger an immediate OCR capture
       - Resize via large corners + edge strips
       - Blocks clicks (no click-through)
       - Center is painted with alpha=1 so macOS hit-tests it
       - Hosts the AnswerCard along its bottom edge (masked out of captures)
    """
    EDGE_T = 12
    CORNER = 26
    MIN_W  = 240   # small enough for a one-line answer pill, never absurd
    MIN_H  = 120
    HIT_ALPHA = 1  # barely visible fill so the center is hit-testable
    OCR_CLICK_RADIUS = 40  # px radius around center that counts as "click" for OCR
    CARD_MARGIN = 12

    adjusting = pyqtSignal(bool)
    geometry_changed = pyqtSignal(QRect)
    pauseRequested = pyqtSignal(bool)  # True=start drag -> pause, False=end drag -> resume (debounced by header)
    ocrRequested = pyqtSignal()
    card_visibility_changed = pyqtSignal(bool)

    def __init__(self, geo: QRect, corner_radius: int = 22):
        super().__init__(
            None,
            Qt.WindowType.FramelessWindowHint
            | Qt.WindowType.WindowStaysOnTopHint
        )
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground, True)
        self.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents, False)  # not click-through
        self.setStyleSheet("background: transparent;")
        self.setGeometry(geo)
        self._corner_radius = corner_radius

        # drag bookkeeping
        self._drag_origin = None
        self._press_global = QPoint()
        self._press_pos = QPoint()
        self._dragging = False  # track if a move actually happened

        # grip state
        self._grip_active = None
        self._grip_origin_geo = None
        self._grip_press_global = QPoint()

        # robust adjusting flag
        self._is_adjusting = False

        # answer card (inside the pane, masked from OCR)
        self.card = AnswerCard(self)
        self.card.layout_changed.connect(self._reflow)
        self.card.dock_toggle_requested.connect(self.flip_card_side)
        self._card_side = load_config().get("card_side", "bottom")
        if self._card_side not in ("top", "bottom"):
            self._card_side = "bottom"
        self._card_blur = None
        self._card_was_visible = None

        # build hit areas
        self._make_perimeter_bands()
        self._make_resize_grips()
        self._reflow()

        # ensure the whole rect is hit-testable
        self.setMask(QRegion(self.rect()))

        # glow pulse
        self._glow_level = 0.0  # 0..1 fade
        self._glow_timer = QTimer(self)
        self._glow_timer.setInterval(16)  # ~60 FPS
        self._glow_timer.timeout.connect(self._advance_glow)
        self._glow_decay_ms = 650  # total glow duration (longer so it's visible)
        self._glow_started_at = 0.0

    # ----- masking API (used by Engine) -----
    def mask_rects(self) -> List[QRect]:
        """Pane-local logical rects that must be blanked before OCR."""
        if self.card.isHidden():
            return []
        return [QRect(self.card.geometry())]

    # ----- card dock side -----
    @property
    def card_side(self) -> str:
        return self._card_side

    def set_card_side(self, side: str):
        if side not in ("top", "bottom") or side == self._card_side:
            return
        self._card_side = side
        cfg = load_config(); cfg["card_side"] = side
        try:
            save_config(cfg)
        except OSError:
            pass
        self._reflow()

    def flip_card_side(self):
        self.set_card_side("top" if self._card_side == "bottom" else "bottom")

    # ----- color API -----
    def set_color(self, color: QColor):
        global PANE_COLOR
        PANE_COLOR = QColor(color)
        self._style_bands_and_grips()
        self.update()

    # ----- glow API -----
    def trigger_glow(self):
        self._glow_started_at = time.time()
        self._glow_level = 1.0
        if not self._glow_timer.isActive():
            self._glow_timer.start()
        self.update()

    def _advance_glow(self):
        elapsed = (time.time() - self._glow_started_at) * 1000.0
        t = max(0.0, 1.0 - (elapsed / self._glow_decay_ms))  # linear fade 1→0
        self._glow_level = t
        if t <= 0.0:
            self._glow_timer.stop()
            self._glow_level = 0.0
        self.update()

    # ----- adjusting helpers -----
    def _begin_adjusting(self):
        if not self._is_adjusting:
            self._is_adjusting = True
            self.adjusting.emit(True)

    def _end_adjusting(self):
        if self._is_adjusting:
            self._is_adjusting = False
            self.adjusting.emit(False)

    # ---------- center-drag the whole pane ----------
    def mousePressEvent(self, e):
        if e.button() == Qt.MouseButton.LeftButton:
            self._drag_origin = self.geometry()
            self._press_global = e.globalPosition().toPoint()
            self._press_pos = e.position().toPoint()
            self._dragging = False  # reset; decide on click vs drag on move

    def mouseMoveEvent(self, e):
        if e.buttons() & Qt.MouseButton.LeftButton and self._drag_origin is not None:
            if not self._dragging:
                self._dragging = True
                self._begin_adjusting()
                self.pauseRequested.emit(True)  # << pause now # ensure we pause even if press was missed
            delta = e.globalPosition().toPoint() - self._press_global
            self.setGeometry(self._drag_origin.translated(delta))
            self.geometry_changed.emit(self.geometry())

    def mouseReleaseEvent(self, e):
        if e.button() == Qt.MouseButton.LeftButton:
            if self._dragging:
                self._drag_origin = None
                self._dragging = False
                self._end_adjusting()
                self.geometry_changed.emit(self.geometry())
                self.pauseRequested.emit(False)  # << allow resume (header debounces)
            else:
                self._drag_origin = None
                center = self.rect().center()
                if (self._press_pos - center).manhattanLength() <= self.OCR_CLICK_RADIUS:
                    self.ocrRequested.emit()
            self._press_pos = QPoint()

    # ---------- perimeter move bands ----------
    def _make_perimeter_bands(self):
        self.move_bands = {}
        for side in ("top", "bottom", "left", "right"):
            band = QWidget(self)
            band.setObjectName(f"move_{side}")
            band.setCursor(Qt.CursorShape.SizeAllCursor)
            band.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents, False)
            band.mousePressEvent = self._band_press
            band.mouseMoveEvent  = self._band_move
            band.mouseReleaseEvent = self._band_release
            self.move_bands[side] = band
        self._style_bands_and_grips()

    def _style_bands_and_grips(self):
        c = PANE_COLOR
        for band in getattr(self, "move_bands", {}).values():
            band.setStyleSheet(f"background: rgba({c.red()},{c.green()},{c.blue()},0.12);")
        for g in getattr(self, "grips", {}).values():
            g.setStyleSheet(
                f"""
                QWidget#grip_nw, QWidget#grip_ne, QWidget#grip_sw, QWidget#grip_se {{
                    background: transparent;
                }}
                QWidget#grip_n:hover, QWidget#grip_s:hover, QWidget#grip_w:hover, QWidget#grip_e:hover,
                 QWidget#grip_nw:hover, QWidget#grip_ne:hover, QWidget#grip_sw:hover, QWidget#grip_se:hover {{
                    background: rgba({c.red()},{c.green()},{c.blue()},0.20);
                    border-radius: 6px;
               }}
            """
            )

    def _band_press(self, e):
        if e.button() == Qt.MouseButton.LeftButton:
            self._drag_origin = self.geometry()
            self._press_global = e.globalPosition().toPoint()
            self._begin_adjusting()
            self.pauseRequested.emit(True)  # << pause now

    def _band_move(self, e):
        if e.buttons() & Qt.MouseButton.LeftButton and self._drag_origin is not None:
            self._begin_adjusting()
            delta = e.globalPosition().toPoint() - self._press_global
            self.setGeometry(self._drag_origin.translated(delta))
            self.geometry_changed.emit(self.geometry())

    def _band_release(self, e):
        if e.button() == Qt.MouseButton.LeftButton:
            self._drag_origin = None
            self._end_adjusting()
            self.geometry_changed.emit(self.geometry())
            self.pauseRequested.emit(False)  # << allow resume

    # ---------- resize grips ----------
    def _make_resize_grips(self):
        self.grips = {}
        def mk(kind: str, cursor: Qt.CursorShape):
            g = QWidget(self)
            g.setObjectName(f"grip_{kind}")
            g.setCursor(cursor)
            g.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents, False)
            g.mousePressEvent = lambda e, k=kind: self._grip_press(e, k)
            g.mouseMoveEvent  = lambda e, k=kind: self._grip_move(e, k)
            g.mouseReleaseEvent = lambda e, k=kind: self._grip_release(e, k)
            self.grips[kind] = g

        mk("nw", Qt.CursorShape.SizeFDiagCursor)
        mk("ne", Qt.CursorShape.SizeBDiagCursor)
        mk("sw", Qt.CursorShape.SizeBDiagCursor)
        mk("se", Qt.CursorShape.SizeFDiagCursor)
        mk("n",  Qt.CursorShape.SizeVerCursor)
        mk("s",  Qt.CursorShape.SizeVerCursor)
        mk("w",  Qt.CursorShape.SizeHorCursor)
        mk("e",  Qt.CursorShape.SizeHorCursor)
        self._style_bands_and_grips()

    def _grip_press(self, e, kind):
        if e.button() != Qt.MouseButton.LeftButton:
            return
        self._grip_active = kind
        self._grip_origin_geo = self.geometry()
        self._grip_press_global = e.globalPosition().toPoint()
        self._begin_adjusting()
        self.pauseRequested.emit(True)  # << pause now

    def _grip_move(self, e, kind):
        if (e.buttons() & Qt.MouseButton.LeftButton) == 0 or self._grip_active != kind:
            return
        self._begin_adjusting()

        delta = e.globalPosition().toPoint() - self._grip_press_global
        x = self._grip_origin_geo.x(); y = self._grip_origin_geo.y()
        w = self._grip_origin_geo.width(); h = self._grip_origin_geo.height()
        dx, dy = delta.x(), delta.y()

        if "w" in kind:
            nx, nw = x + dx, w - dx
            if nw >= self.MIN_W: x, w = nx, nw
        if "e" in kind:
            nw = w + dx
            if nw >= self.MIN_W: w = nw
        if "n" in kind:
            ny, nh = y + dy, h - dy
            if nh >= self.MIN_H: y, h = ny, nh
        if "s" in kind:
            nh = h + dy
            if nh >= self.MIN_H: h = nh

        self.setGeometry(QRect(x, y, w, h))
        self.geometry_changed.emit(self.geometry())

    def _grip_release(self, e, kind):
        if e.button() != Qt.MouseButton.LeftButton:
            return
        self._grip_active = None
        self._grip_origin_geo = None
        self._end_adjusting()
        self.geometry_changed.emit(self.geometry())
        self.pauseRequested.emit(False)  # << allow resume

    # ---------- layout ----------
    def _layout_perimeter_bands(self):
        r = self.rect(); t = self.EDGE_T; g = self.CORNER
        # inset bands so they don't overlap corner grips
        self.move_bands["top"].setGeometry(g, 0, max(0, r.width() - 2*g), t)
        self.move_bands["bottom"].setGeometry(g, r.height()-t, max(0, r.width() - 2*g), t)
        self.move_bands["left"].setGeometry(0, g, t, max(0, r.height() - 2*g))
        self.move_bands["right"].setGeometry(r.width()-t, g, t, max(0, r.height() - 2*g))

    def _layout_grips(self):
        r = self.rect(); g = self.CORNER; t = self.EDGE_T
        self.grips["nw"].setGeometry(0, 0, g, g)
        self.grips["ne"].setGeometry(r.width()-g, 0, g, g)
        self.grips["sw"].setGeometry(0, r.height()-g, g, g)
        self.grips["se"].setGeometry(r.width()-g, r.height()-g, g, g)
        self.grips["n"].setGeometry(g, 0, r.width()-2*g, t)
        self.grips["s"].setGeometry(g, r.height()-t, r.width()-2*g, t)
        self.grips["w"].setGeometry(0, g, t, r.height()-2*g)
        self.grips["e"].setGeometry(r.width()-t, g, t, r.height()-2*g)
        for k in ("nw", "ne", "sw", "se"):
            self.grips[k].raise_()
        for k in ("n", "s", "w", "e"):
            self.grips[k].raise_()

    def _layout_card(self):
        m = self.CARD_MARGIN
        avail_w = self.width() - 2 * m
        tier = AnswerCard.tier_for(self.width(), self.height())
        self.card.set_tier(tier, self.width())
        show = tier != "hidden"
        if show:
            h = self.card.desired_height(self.height())
            h = min(h, self.height() - 2 * m - self.EDGE_T)
            y = m if self._card_side == "top" else self.height() - m - h
            self.card.setGeometry(m, y, avail_w, h)
            self.card.show()
            self.card.raise_()
        else:
            self.card.hide()
        reframe_native_blur(self, self._card_blur, self.card.geometry(), visible=show)
        if show != self._card_was_visible:
            self._card_was_visible = show
            self.card_visibility_changed.emit(show)

    def _reflow(self):
        if not hasattr(self, "move_bands") or not hasattr(self, "grips"):
            return
        self._layout_perimeter_bands()
        self._layout_grips()
        self._layout_card()

    def resizeEvent(self, _):
        self._reflow()
        self.setMask(QRegion(self.rect()))

    def showEvent(self, e):
        super().showEvent(e)
        if self._card_blur is None:
            QTimer.singleShot(0, self._attach_card_blur)

    def _attach_card_blur(self):
        if self._card_blur is None:
            self._card_blur = apply_native_blur(self, self.card.radius, rect=self.card.geometry())
            self.card.set_blur_active(self._card_blur is not None)
            reframe_native_blur(self, self._card_blur, self.card.geometry(), visible=not self.card.isHidden())

    # ---------- paint ----------
    def paintEvent(self, _):
        p = QPainter(self)
        p.setRenderHint(QPainter.RenderHint.Antialiasing, True)

        # fill center with 1-alpha for hit-testing (interior must stay clear for OCR)
        p.fillRect(self.rect(), QColor(0, 0, 0, self.HIT_ALPHA))

        r = QRectF(self.rect()).adjusted(2, 2, -2, -2)

        # --- yellowish-white pulse when active (additive so it "blooms") ---
        if self._glow_level > 0.0:
            glow_width = 14 + int(14 * self._glow_level)
            glow_alpha = int(220 * self._glow_level)
            pulse_color = QColor(255, 248, 210, glow_alpha)
            p.save()
            p.setCompositionMode(QPainter.CompositionMode.CompositionMode_Plus)
            pen = QPen(pulse_color, glow_width,
                       cap=Qt.PenCapStyle.RoundCap,
                       join=Qt.PenJoinStyle.RoundJoin)
            p.setPen(pen)
            p.setBrush(Qt.BrushStyle.NoBrush)
            p.drawRoundedRect(r, self._corner_radius, self._corner_radius)
            p.restore()

        # --- colored glass frame: soft outer stroke + crisp inner highlight ---
        p.setBrush(Qt.BrushStyle.NoBrush)
        p.setPen(QPen(PANE_COLOR, 3.5))
        p.drawRoundedRect(r, self._corner_radius, self._corner_radius)
        hg = QLinearGradient(r.topLeft(), r.bottomRight())
        hg.setColorAt(0.0, QColor(255, 255, 255, 170))
        hg.setColorAt(0.6, QColor(255, 255, 255, 30))
        p.setPen(QPen(QBrush(hg), 1.0))
        p.drawRoundedRect(r.adjusted(2.5, 2.5, -2.5, -2.5), self._corner_radius - 2, self._corner_radius - 2)


class Engine(QObject):
    answer_ready = pyqtSignal(str, str)  # (context_black, answer_green)
    status_changed = pyqtSignal(str)
    _req_pause = pyqtSignal(str)
    _req_resume = pyqtSignal(str)
    _req_set_bbox = pyqtSignal(object)
    _req_set_interval = pyqtSignal(int)
    ocr_pulse = pyqtSignal()

    def __init__(self, bbox: CaptureRect,
                 min_interval_ms: int = None,
                 max_interval_ms: int = None,
                 mask_provider: Optional[Callable[[], List[QRect]]] = None):
        super().__init__()
        # pacing
        self._min_interval = int(os.getenv("MM_MIN_INTERVAL_MS", min_interval_ms or 1500))
        self._max_interval = int(os.getenv("MM_MAX_INTERVAL_MS", max_interval_ms or 5000))
        self._interval = self._min_interval

        self._bbox = bbox
        self._mask_provider = mask_provider
        self._timer = QTimer(self)
        self._timer.setInterval(self._interval)
        self._timer.timeout.connect(self._tick)

        self._busy = False
        self._last_text = ""
        self._paused = False
        self.set_status = lambda s: self.status_changed.emit(s)

        # queued connections ensure these slots run in the Qt (main) thread
        self._req_pause.connect(self._do_pause, Qt.ConnectionType.QueuedConnection)
        self._req_resume.connect(self._do_resume, Qt.ConnectionType.QueuedConnection)
        self._req_set_bbox.connect(self._do_set_bbox, Qt.ConnectionType.QueuedConnection)
        self._req_set_interval.connect(self._do_set_interval, Qt.ConnectionType.QueuedConnection)
        self._gen = 0  # increments on every pause to invalidate in-flight work

    # ---- capture loop ----
    def _tick(self):
        if self._busy or self._paused:
            return
        self._busy = True
        try:
            self.set_status(f"capturing… ({self._interval} ms)")
            screen = QGuiApplication.primaryScreen()
            pix = screen.grabWindow(0, self._bbox.left, self._bbox.top, self._bbox.width, self._bbox.height)
            qimg = pix.toImage()

            # blank the on-screen answer card out of the capture BEFORE anything reads it
            if self._mask_provider is not None:
                mask_image(qimg, self._mask_provider())

            # 🔔 fire the pulse right as an OCR cycle begins
            self.ocr_pulse.emit()

            # capture current generation for cancellation
            gen = self._gen
        except Exception as e:
            self.answer_ready.emit("", f"(capture error: {e})")
            self.set_status("capture error")
            self._busy = False
            return

        # run OCR/LLM in background, with gen for cancelation
        threading.Thread(target=self._ocr_and_ask, args=(qimg, gen), daemon=True).start()

    def start(self):
        if not self._paused:
            self._timer.start()

    def stop(self):
        self._timer.stop()

    def pause(self, reason="paused"):
        self._req_pause.emit(reason)

    def resume(self, reason="resumed"):
        self._req_resume.emit(reason)

    def set_bbox(self, bbox: CaptureRect):
        self._req_set_bbox.emit(bbox)

    def capture_now(self):
        """Trigger an immediate capture/OCR cycle if idle."""
        if not self._busy and not self._paused:
            self._tick()

    def set_interval_ms(self, ms: int):
        self._req_set_interval.emit(ms)

    @property
    def paused(self) -> bool:
        return self._paused

    # ---- slots (main thread only) ----
    @pyqtSlot(str)
    def _do_pause(self, reason):
        self._paused = True
        self._gen += 1  # << invalidate any in-flight OCR/LLM
        self._timer.stop()
        self.set_status(f"{reason} ({self._interval} ms)")

    @pyqtSlot(str)
    def _do_resume(self, reason):
        self._paused = False
        self._timer.start()
        self.set_status(f"{reason} ({self._interval} ms)")

    @pyqtSlot(object)
    def _do_set_bbox(self, bbox: CaptureRect):
        self._bbox = bbox
        self._maybe_reset_to_min("⏸️ adjusted 📏")

    @pyqtSlot(int)
    def _do_set_interval(self, ms: int):
        self._interval = max(50, min(int(ms), self._max_interval))
        self._timer.setInterval(self._interval)
        self.set_status(f"⏸️⏲️ interval {self._interval} ms")

    # ---- pacing helpers (thread-safe wrappers) ----
    def _maybe_backoff(self, why: str):
        new_int = int(min(self._max_interval, max(self._interval, int(self._interval * 1.6))))
        if new_int != self._interval:
            self._req_set_interval.emit(new_int)
        self.set_status(f"{why} — backoff to {max(new_int, self._interval)} ms")

    def _maybe_reset_to_min(self, why: str):
        if self._interval != self._min_interval:
            self._req_set_interval.emit(self._min_interval)
            self.set_status(f"{why} — reset to {self._min_interval} ms")

    def _ocr_and_ask(self, qimg: QImage, gen: int):
        # cancel if a pause happened after this job started
        if gen != self._gen:
            self._busy = False
            return

        self.set_status(f"OCR ({OCR_BACKEND})…")

        # ---- OCR
        try:
            text = run_ocr(qimg)
        except Exception as e:
            if gen == self._gen:
                self.answer_ready.emit("", f"(OCR error: {e})")
                self.set_status("OCR error")
            self._busy = False
            return

        # cancel after OCR
        if gen != self._gen:
            self._busy = False
            return

        # pacing / empty text handling
        if not text:
            self.answer_ready.emit("", "")
            self.set_status("no text")
            self._last_text = ""
            self._maybe_backoff("no text")
            self._busy = False
            return

        if text == self._last_text:
            self._maybe_backoff("no change")
            self._busy = False
            return

        self._maybe_reset_to_min("new text")
        self.set_status(f"OCR… ({len(text)} chars)")
        self._last_text = text

        # ---- LLM call
        answer = ""  # <-- ensure defined even if something fails below
        if gen != self._gen:
            self._busy = False
            return

        try:
            self.set_status("asking…")
            messages = build_messages(text)
            print("\n[MirrorMirror] Sending to ChatGPT model:")
            print("Model:", MODEL)
            print("System:\n", messages[0]["content"][:2000], "…" if len(messages[0]["content"]) > 2000 else "")
            print("User:\n", messages[1]["content"])
            sys.stdout.flush()

            resp = get_client().chat.completions.create(
                model=MODEL,
                messages=messages,
                reasoning_effort=REASONING_EFFORT,
            )
            answer = resp.choices[0].message.content.strip()
            self.set_status("✅ Answer Received!")
        except Exception as e:
            answer = f"(OpenAI error: {e})"
            self.set_status("LLM error")

        # cancel after LLM if we were paused mid-flight
        if gen != self._gen:
            self._busy = False
            return

        # final emit
        self.answer_ready.emit("", answer)
        self._busy = False


def build_messages(screen_text: str) -> List[dict]:
    """System message (instructions + reference material, static so it caches) and
    the user message with the fresh OCR text."""
    block, names = get_refs().context_for(screen_text)
    system = (
        "You are a research assistant reading text captured from a screen by OCR "
        "(it may contain typos or stray characters).\n"
        "Scan the screen text for a question. If there is no question, reply with one short helpful sentence.\n"
        "If it looks like a multiple-choice question, answer with the correct option.\n"
        "Rules:\n"
        "1. Do not guess. If the answer is uncertain, say so.\n"
        "2. Draw on reliable knowledge (academic, government, encyclopedic, technical references).\n"
        "3. If sources conflict, note both sides and the consensus.\n"
        "4. Keep it accurate, neutral, and free of speculation.\n"
        "5. Do not restate or summarize the question.\n"
        "Format your reply EXACTLY as:\n"
        "Line 1: the answer in 3 words or fewer when possible (for multiple choice, the option letter and text).\n"
        "Line 2: Confidence: NN%\n"
    )
    if block:
        system += (
            "Line 3: Source: <file name from the reference material, or 'general knowledge'>\n"
            "Then, optionally, one or two short sentences of explanation.\n\n"
            "REFERENCE MATERIAL: the user supplied the files below. Treat them as the primary "
            "source. If they cover the question, base the answer on them (even if it differs "
            "from what you would otherwise say) and name the file on the Source line. Only if "
            "they do not cover the question fall back to general knowledge and write "
            "'Source: general knowledge'.\n\n"
            "=== REFERENCE MATERIAL ===\n" + block + "\n=== END REFERENCE MATERIAL ==="
        )
    else:
        system += "Then, optionally, one or two short sentences of explanation.\n"
    user = f"Screen text (OCR):\n---\n{screen_text}\n---"
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


# ================== References window ==================
class ReferencesWindow(GlassWindow):
    """Manage the files the model should prefer when answering."""
    KIND_ICON = {"pdf": "📕", "word": "📘", "slides": "📙", "image": "🖼️", "text": "📄"}

    def __init__(self):
        super().__init__("References", closable=True, min_width=520)
        self.setAcceptDrops(True)
        refs = get_refs()

        intro = QLabel(
            "Add notes, textbooks, slides or photos. They are read on this Mac and sent to the "
            "model with every question, and the answer will prefer this material over general knowledge."
        )
        intro.setWordWrap(True)
        self.body.addWidget(intro)

        top = QHBoxLayout(); top.setSpacing(8)
        self._use_chk = QCheckBox("Use references when answering")
        self._use_chk.setChecked(refs.enabled)
        self._use_chk.toggled.connect(refs.set_enabled)
        top.addWidget(self._use_chk)
        top.addStretch(1)
        self.body.addLayout(top)

        self._scroll = QScrollArea()
        self._scroll.setWidgetResizable(True)
        self._scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self._scroll.setFrameShape(QFrame.Shape.NoFrame)
        self._scroll.setMinimumHeight(160)
        self._scroll.setMaximumHeight(320)
        self._host = QWidget()
        self._list = QVBoxLayout(self._host)
        self._list.setContentsMargins(0, 0, 4, 0)
        self._list.setSpacing(6)
        self._list.addStretch(1)
        self._scroll.setWidget(self._host)
        self.body.addWidget(self._scroll)

        self._empty = QLabel("No references yet. Drop files here or click Add Files.")
        self._empty.setObjectName("subtext")
        self._empty.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.body.addWidget(self._empty)

        self._footer = QLabel("")
        self._footer.setObjectName("subtext")
        self._footer.setWordWrap(True)
        self.body.addWidget(self._footer)

        btns = QHBoxLayout(); btns.setSpacing(8)
        add_btn = QPushButton("Add Files…"); add_btn.setObjectName("primary")
        add_btn.clicked.connect(self._add_files)
        self._clear_btn = QPushButton("Clear All")
        self._clear_btn.clicked.connect(refs.clear)
        done_btn = QPushButton("Done")
        done_btn.clicked.connect(self.close)
        btns.addWidget(add_btn)
        btns.addWidget(self._clear_btn)
        btns.addStretch(1)
        btns.addWidget(done_btn)
        self.body.addLayout(btns)

        refs.changed.connect(self._rebuild)
        self._rebuild()
        self.adjustSize()

    def _add_files(self):
        # parent=None so the native dialog doesn't inherit the glass stylesheet
        paths, _ = QFileDialog.getOpenFileNames(None, "Add reference files", str(Path.home()), REF_FILE_FILTER)
        if paths:
            get_refs().add_paths(paths)

    # drag & drop of files onto the window
    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls() and any(u.isLocalFile() for u in e.mimeData().urls()):
            e.acceptProposedAction()

    def dropEvent(self, e):
        paths = [u.toLocalFile() for u in e.mimeData().urls() if u.isLocalFile()]
        if paths:
            get_refs().add_paths(paths)
            e.acceptProposedAction()

    def _rebuild(self):
        refs = get_refs()
        while self._list.count() > 1:
            item = self._list.takeAt(0)
            if item and item.widget():
                item.widget().deleteLater()
        docs = refs.docs
        for d in docs:
            self._list.insertWidget(self._list.count() - 1, self._row(d))
        self._empty.setVisible(not docs)
        self._scroll.setVisible(bool(docs))
        self._clear_btn.setEnabled(bool(docs))
        n = len(docs); ready = refs.ready_docs(); chars = refs.total_chars()
        if not docs:
            self._footer.setText("")
        else:
            pending = sum(1 for d in docs if d.status == "pending")
            how = ("everything fits in one prompt and is cached" if refs.fits_in_one_prompt()
                   else "larger than one prompt, so only the most relevant excerpts are sent each time")
            parts = [f"{n} file{'s' if n != 1 else ''}", f"{chars/1000:.1f}k characters ready"]
            if pending:
                parts.append(f"{pending} indexing…")
            self._footer.setText(" · ".join(parts) + f" — {how}.")
        self._use_chk.setChecked(refs.enabled)

    def _row(self, d: RefDoc) -> QWidget:
        row = QFrame(); row.setObjectName("row")
        h = QHBoxLayout(row); h.setContentsMargins(10, 6, 8, 6); h.setSpacing(8)
        icon = QLabel(self.KIND_ICON.get(d.kind, "📄")); icon.setFixedWidth(22)
        name = QLabel(d.name); name.setObjectName("title")
        name.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        name.setToolTip(d.path)
        status = {
            "ready": f"{d.chars/1000:.1f}k chars",
            "pending": "indexing…",
            "empty": "no text found",
            "error": f"error: {d.error}",
        }.get(d.status, d.status)
        st = QLabel(status); st.setObjectName("error" if d.status in ("error", "empty") else "subtext")
        rm = QPushButton("✕"); rm.setObjectName("ghost"); rm.setFixedSize(24, 24); rm.setToolTip("Remove")
        rm.clicked.connect(lambda _=False, i=d.id: get_refs().remove(i))
        h.addWidget(icon); h.addWidget(name, 1); h.addWidget(st); h.addWidget(rm)
        return row


# ================== Header (toolbar) window ==================
class HeaderWindow(GlassWindow):
    RADIUS = 20
    PANE_SIZE = (800, 400)

    def __init__(self):
        super().__init__("Mirror Mirror", closable=True, min_width=420)
        self.close_btn.setToolTip("Quit")

        # status dot before the title
        self._dot = QLabel(); self._dot.setObjectName("dot"); self._dot.setFixedSize(10, 10)
        self.title_row.insertWidget(0, self._dot)

        # toolbar row
        row = QHBoxLayout(); row.setSpacing(8)
        self._pause_btn = QPushButton("Pause")
        self._pause_btn.clicked.connect(self._toggle_user_pause)
        clear_btn = QPushButton("Clear")
        clear_btn.clicked.connect(self._clear_answers)
        self._refs_btn = QPushButton("References")
        self._refs_btn.setToolTip("Add files the model should prefer when answering")
        self._refs_btn.clicked.connect(self._open_references)
        self._refs_win = None
        get_refs().changed.connect(self._update_refs_btn)
        self._update_refs_btn()
        settings_btn = QPushButton("Settings")
        settings_btn.clicked.connect(self._open_settings)
        quit_btn = QPushButton("Quit"); quit_btn.setObjectName("danger")
        quit_btn.clicked.connect(QApplication.instance().quit)
        row.addWidget(self._pause_btn)
        row.addWidget(clear_btn)
        row.addWidget(self._refs_btn)
        row.addStretch(1)
        row.addWidget(settings_btn)
        row.addWidget(quit_btn)
        self.body.addLayout(row)

        hint = QLabel("Drag the gold pane over text. Click its center to read now.")
        hint.setObjectName("subtext")
        self.body.addWidget(hint)

        # fallback answer row, only when the pane is too small to host the card
        self._latest_row = QWidget()
        lr = QHBoxLayout(self._latest_row)
        lr.setContentsMargins(0, 0, 0, 0); lr.setSpacing(8)
        self._latest_badge = ConfidenceBadge(show_text=True)
        self._latest_lbl = QLabel("")
        self._latest_lbl.setObjectName("headline_small")
        self._latest_lbl.setWordWrap(True)
        lr.addWidget(self._latest_badge, 0, Qt.AlignmentFlag.AlignTop)
        lr.addWidget(self._latest_lbl, 1)
        self._latest_row.hide()
        self.body.addWidget(self._latest_row)

        self.status_lbl = QLabel("idle")
        self.status_lbl.setObjectName("subtext")
        self.status_lbl.setVisible(SHOW_DEBUG_STATUS)
        self.body.addWidget(self.status_lbl)

        self.adjustSize()
        self.move(140, 120)

        # pause debounce timer
        self._pause_timer = QTimer(self)
        self._pause_timer.setSingleShot(True)
        self._pause_timer.timeout.connect(self._resume_engine_after_pause)
        self._user_paused = False
        self._error_timer = QTimer(self)
        self._error_timer.setSingleShot(True)
        self._error_timer.timeout.connect(lambda: self._set_dot("ok"))

        self._glass = None
        self._engine = None
        self._settings = None

        # 1) glass pane
        screen = QGuiApplication.primaryScreen()
        geo = screen.geometry()
        pw, ph = self.PANE_SIZE
        bbox = CaptureRect((geo.width() - pw) // 2, (geo.height() - ph) // 2, pw, ph)
        self._glass = GlassPane(QRect(bbox.left, bbox.top, bbox.width, bbox.height))
        self._glass.show()

        # 2) engine (masks the answer card out of every capture)
        self._engine = Engine(bbox, mask_provider=self._glass.mask_rects)

        # 3) engine -> UI
        self._engine.status_changed.connect(self.set_status, type=Qt.ConnectionType.QueuedConnection)
        self._engine.answer_ready.connect(self._on_answer, type=Qt.ConnectionType.QueuedConnection)
        self._engine.ocr_pulse.connect(self._glass.trigger_glow, type=Qt.ConnectionType.QueuedConnection)

        # 4) glass -> engine/UI
        self._glass.pauseRequested.connect(self._on_pause_request, type=Qt.ConnectionType.QueuedConnection)
        self._glass.adjusting.connect(self._on_glass_adjusting, type=Qt.ConnectionType.QueuedConnection)
        self._glass.geometry_changed.connect(self._on_glass_geometry_changed, type=Qt.ConnectionType.QueuedConnection)
        self._glass.ocrRequested.connect(self._engine.capture_now, type=Qt.ConnectionType.QueuedConnection)
        self._glass.card_visibility_changed.connect(self._on_card_visibility)

        self._set_dot("ok")
        self.set_status(f"ready (ocr: {OCR_BACKEND})")
        self._engine.start()

    # ----- status dot -----
    def _set_dot(self, state: str):
        color = {"ok": "#30d158", "paused": "#ffd60a", "error": "#ff453a"}.get(state, "#8e8e93")
        self._dot.setStyleSheet(f"background: {color}; border-radius: 5px;")

    def set_status(self, msg: str):
        if SHOW_DEBUG_STATUS:
            self.status_lbl.setText(msg or "")
            self.status_lbl.show()
        else:
            self.status_lbl.hide()
        if self._engine and self._engine.paused:
            self._set_dot("paused")
        elif not self._error_timer.isActive():
            self._set_dot("ok")

    # ----- answers -----
    def _on_answer(self, black: str, green: str):
        if not green:
            return
        card = self._glass.card if self._glass else None
        if _is_error_text(green):
            if card:
                card.set_error(green)
            self._set_dot("error")
            self._error_timer.start(4000)
            return
        if card:
            card.push(green)
            self._update_latest_row()

    def _clear_answers(self):
        if self._glass:
            self._glass.card.clear()
        self._update_latest_row()

    def _update_latest_row(self):
        card = self._glass.card if self._glass else None
        a = card.latest if card else None
        self._latest_lbl.setText(a.headline if a else "")
        set_level(self._latest_lbl, confidence_level(a.conf_value) if a else "unknown")
        self._latest_badge.set_answer(a)
        self._latest_row.setVisible(bool(a) and bool(card) and card.isHidden())

    def _on_card_visibility(self, visible: bool):
        self._update_latest_row()

    # ----- pause / resume -----
    def _toggle_user_pause(self):
        if not self._engine:
            return
        self._user_paused = not self._user_paused
        if self._user_paused:
            self._pause_timer.stop()
            self._engine.pause("⏸️ paused")
            self._pause_btn.setText("Resume")
            self._pause_btn.setObjectName("primary")
        else:
            self._engine.resume("▶️ resumed")
            self._pause_btn.setText("Pause")
            self._pause_btn.setObjectName("")
        # re-polish so the objectName change picks up the new style
        self._pause_btn.style().unpolish(self._pause_btn)
        self._pause_btn.style().polish(self._pause_btn)
        QTimer.singleShot(0, lambda: self._set_dot("paused" if self._user_paused else "ok"))

    def _on_pause_request(self, should_pause: bool):
        if not self._engine:
            return
        if should_pause:
            self._engine.pause("⏸️ paused (dragging ✊)")
            self._pause_timer.stop()
        else:
            self._pause_timer.start(350)  # resume shortly after final release (debounced)

    def _resume_engine_after_pause(self):
        if self._engine and not self._user_paused:
            self._engine.resume("▶️ resumed")

    # pause/resume around header moves/resizes
    def moveEvent(self, e):
        super().moveEvent(e)
        if self._engine:
            self._engine.pause("⏸️ paused (moving header 📏)")
            self._pause_timer.start(350)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        if self._engine:
            self._engine.pause("⏸️ paused (resizing header 📏)")
            self._pause_timer.start(350)

    # ----- glass hooks -----
    def _on_glass_adjusting(self, is_adjusting: bool):
        if not self._engine: return
        if is_adjusting:
            self._engine.pause("⏸️ paused (adjusting 📏)")
            self._pause_timer.stop()
        else:
            self._pause_timer.start(350)

    def _on_glass_geometry_changed(self, new_geo: QRect):
        if not self._engine:
            return
        self._engine.pause("⏸️ paused (adjusting 📏)")
        self._pause_timer.stop()
        self._pause_timer.start(350)  # resume a bit after the last change
        bbox = CaptureRect(new_geo.x(), new_geo.y(), new_geo.width(), new_geo.height())
        self._engine.set_bbox(bbox)
        self.set_status("adjusted 📏")

    # ----- references -----
    def _update_refs_btn(self):
        refs = get_refs()
        n = len(refs.docs)
        self._refs_btn.setText(f"References ({n})" if n else "References")
        self._refs_btn.setObjectName("primary" if (n and refs.enabled) else "")
        st = self._refs_btn.style(); st.unpolish(self._refs_btn); st.polish(self._refs_btn)

    def _open_references(self):
        if self._refs_win is None or not self._refs_win.isVisible():
            self._refs_win = ReferencesWindow()
            self._refs_win.move(self.x(), self.y() + self.height() + 12)
            self._refs_win.show()
        self._refs_win.raise_()
        self._refs_win.activateWindow()

    # ----- settings / lifecycle -----
    def _teardown(self):
        if self._engine:
            self._engine.stop()
        if self._glass:
            self._glass.close()
        if self._refs_win:
            self._refs_win.close()

    def _open_settings(self):
        """Close header and show settings window."""
        self._teardown()

        def reopen():
            return HeaderWindow()

        self._settings = SettingsWindow(reopen)
        self._settings.show()
        self.close()

    def closeEvent(self, e):
        self._teardown()
        super().closeEvent(e)


# ================== Settings window ==================
class SettingsWindow(GlassWindow):
    """Simple window for adjusting application settings."""

    def __init__(self, return_factory=None):
        super().__init__("Settings", closable=True, min_width=420)

        # --- pane color ---
        sec = QLabel("Capture pane"); sec.setObjectName("subtext")
        self.body.addWidget(sec)
        color_row = QHBoxLayout(); color_row.setSpacing(8)
        self._color_indicator = QLabel()
        self._color_indicator.setFixedSize(24, 24)
        color_row.addWidget(self._color_indicator)
        self._color_btn = QPushButton("Choose Pane Color")
        color_row.addWidget(self._color_btn)
        self._reset_btn = QPushButton("Reset to Default")
        color_row.addWidget(self._reset_btn)
        color_row.addStretch(1)
        self.body.addLayout(color_row)
        self._update_color_display()
        self._color_btn.clicked.connect(self._choose_color)
        self._reset_btn.clicked.connect(self._reset_color)

        self._debug_chk = QCheckBox("Show debug status in toolbar")
        self._debug_chk.setChecked(SHOW_DEBUG_STATUS)
        self.body.addWidget(self._debug_chk)
        self._debug_chk.toggled.connect(self._toggle_debug)

        # --- OpenAI API key ---
        sec2 = QLabel("OpenAI"); sec2.setObjectName("subtext")
        self.body.addWidget(sec2)
        key_row = QHBoxLayout(); key_row.setSpacing(8)
        self._key_lbl = QLabel()
        key_row.addWidget(self._key_lbl)
        key_row.addStretch(1)
        self._key_btn = QPushButton("Change API Key")
        key_row.addWidget(self._key_btn)
        self._key_btn.clicked.connect(self._change_api_key)
        self._update_key_display()
        self.body.addLayout(key_row)

        model_lbl = QLabel(f"Model: {MODEL} · effort: {REASONING_EFFORT} · OCR: {OCR_BACKEND} · v{APP_VERSION}")
        model_lbl.setObjectName("subtext")
        self.body.addWidget(model_lbl)

        back_row = QHBoxLayout()
        back_row.addStretch(1)
        back_btn = QPushButton("Back"); back_btn.setObjectName("primary")
        back_row.addWidget(back_btn)
        self.body.addLayout(back_row)

        self._return_factory = return_factory
        self._return_widget = None
        back_btn.clicked.connect(self._go_back)
        self.adjustSize()

    def _update_color_display(self):
        self._color_indicator.setStyleSheet(
            f"border-radius: 12px; background-color: {PANE_COLOR.name()}; border: 1px solid rgba(0,0,0,0.15);"
        )

    def _choose_color(self):
        global PANE_COLOR
        # parent=None so the native dialog doesn't inherit the glass stylesheet
        color = QColorDialog.getColor(PANE_COLOR, None, "Select Pane Color")
        if color.isValid():
            PANE_COLOR = QColor(color.red(), color.green(), color.blue(), 230)
            self._update_color_display()

    def _reset_color(self):
        global PANE_COLOR
        PANE_COLOR = QColor(DEFAULT_PANE_COLOR)
        self._update_color_display()

    def _toggle_debug(self, checked: bool):
        global SHOW_DEBUG_STATUS
        SHOW_DEBUG_STATUS = checked

    def _update_key_display(self):
        self._key_lbl.setText(f"API key: {mask_key(get_api_key())}")

    def _change_api_key(self):
        dlg = ApiKeyDialog(self, current_key="")
        if dlg.exec() == QDialog.DialogCode.Accepted:
            self._update_key_display()

    def _go_back(self):
        if self._return_factory:
            self._return_widget = self._return_factory()
            self._return_widget.show()
        self.close()


# ================== Welcome window ==================
class WelcomeWindow(GlassWindow):
    """Initial window presenting a friendly welcome and basic instructions."""

    def __init__(self):
        super().__init__("Mirror Mirror", closable=True, min_width=440)
        self.close_btn.setToolTip("Quit")

        lbl = QLabel(
            "Point the gold pane at any text on your screen. Mirror Mirror reads it, "
            "asks OpenAI, and shows a short answer with a confidence score right inside the pane."
        )
        lbl.setWordWrap(True)
        self.body.addWidget(lbl)

        tips = QLabel("Drag the pane's edges to resize · Click its center to read now · Double-click the answer card to collapse it")
        tips.setObjectName("subtext")
        tips.setWordWrap(True)
        self.body.addWidget(tips)

        btn_row = QHBoxLayout(); btn_row.setSpacing(8)
        start_btn = QPushButton("Start"); start_btn.setObjectName("primary")
        settings_btn = QPushButton("Settings")
        quit_btn = QPushButton("Quit")
        btn_row.addStretch(1)
        btn_row.addWidget(quit_btn)
        btn_row.addWidget(settings_btn)
        btn_row.addWidget(start_btn)
        self.body.addLayout(btn_row)

        start_btn.clicked.connect(self._launch)
        settings_btn.clicked.connect(self._open_settings)
        quit_btn.clicked.connect(QApplication.instance().quit)

        self._main = None
        self._settings = None
        self.adjustSize()

    def _launch(self):
        """Close welcome window and show the main HeaderWindow."""
        self._main = HeaderWindow()
        self._main.show()
        self.close()

    def _open_settings(self):
        """Close welcome window and show settings."""
        def reopen():
            return WelcomeWindow()

        self._settings = SettingsWindow(reopen)
        self._settings.show()
        self.close()


# ================== main ==================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)

    # First run (or key removed): ask for a key before showing anything else.
    if not ensure_api_key():
        sys.exit(0)

    w = WelcomeWindow()
    w.show()
    sys.exit(app.exec())
